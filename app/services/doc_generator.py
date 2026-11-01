from __future__ import annotations

import io
import re
import zipfile
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
from docxtpl import DocxTemplate  # type: ignore[import-untyped]
from fastapi import HTTPException, status

from app.models.job_item import JobStatus
from app.models.order import Order, POStatus
from app.models.po_out import PurchaseOrderOut
from app.services.number_generator import (
    format_currency_idr,
    format_date_indo,
    format_day_date_indo,
    format_day_indo,
)

TEMPLATES_DOCS_DIR = Path("app/templates/docs")


# ==============================================================================
# DOCX XML PREPROCESSOR
# ==============================================================================


def merge_tags_in_p(p_xml: str) -> str:
    """Merge split Jinja tags ({{ ... }} and {% ... %}) across multiple <w:r><w:t>

    runs within a single paragraph.
    """
    t_matches = list(re.finditer(r"(<w:t\b[^>]*>)(.*?)(</w:t>)", p_xml, re.DOTALL))
    if not t_matches:
        return p_xml

    full_text = "".join(m.group(2) for m in t_matches)
    if "{{" not in full_text and "{%" not in full_text:
        return p_xml

    char_map: list[tuple[int, int]] = []
    for idx, m in enumerate(t_matches):
        for c_idx in range(len(m.group(2))):
            char_map.append((idx, c_idx))

    tag_matches = list(re.finditer(r"(\{\{.*?\}\}|\{%.*?%\})", full_text, re.DOTALL))
    if not tag_matches:
        return p_xml

    node_chars = [list(m.group(2)) for m in t_matches]

    for tag in tag_matches:
        start_idx = tag.start()
        end_idx = tag.end()

        raw_tag = tag.group(0)
        if raw_tag.startswith("{{") and raw_tag.endswith("}}"):
            inner = raw_tag[2:-2].strip()
            inner = re.sub(r"\s+", " ", inner)
            tag_content = f"{{{{ {inner} }}}}"
        elif raw_tag.startswith("{%") and raw_tag.endswith("%}"):
            inner = raw_tag[2:-2].strip()
            inner = re.sub(r"\s+", " ", inner)
            tag_content = f"{{% {inner} %}}"
        else:
            tag_content = raw_tag

        for c_pos in range(start_idx, end_idx):
            n_idx, c_offset = char_map[c_pos]
            if c_pos == start_idx:
                node_chars[n_idx][c_offset] = tag_content
            else:
                node_chars[n_idx][c_offset] = ""

    res = p_xml
    for idx in reversed(range(len(t_matches))):
        m = t_matches[idx]
        new_content = "".join(node_chars[idx])
        open_tag = m.group(1)
        if (" " in new_content or "\n" in new_content) and "xml:space" not in open_tag:
            open_tag = open_tag.replace("<w:t", '<w:t xml:space="preserve"')
        res = res[: m.start()] + open_tag + new_content + m.group(3) + res[m.end() :]

    return res


def clean_docx_xml_content(xml: str) -> str:
    """Remove MS Word proofing error artifacts, merge split Jinja tags, and wrap

    table rows containing {% tr for ... %} with standard {% for ... %} loops.
    """
    # 1. Remove proofErr tags
    xml = re.sub(r"<w:proofErr[^>]*/>", "", xml)

    # 2. Merge split tags in all paragraphs
    xml = re.sub(
        r"<w:p\b[^>]*>.*?</w:p>",
        lambda m: merge_tags_in_p(m.group(0)),
        xml,
        flags=re.DOTALL,
    )

    # 3. Transform single-row loop tags inside <w:tr>
    def fix_tr(tr_match: re.Match[str]) -> str:
        tr_xml = tr_match.group(0)
        m_start = re.search(r"\{%\s*tr\s+for\s+([^%]+)%\s*\}", tr_xml)
        if not m_start:
            return tr_xml

        loop_var = m_start.group(1).strip()
        tr_clean = re.sub(r"\{%\s*tr\s+for\s+[^%]+%\s*\}", "", tr_xml)
        tr_clean = re.sub(r"\{%\s*(?:endtr|tr\s+endfor)\s*%\s*\}", "", tr_clean)
        tr_clean = re.sub(r"Kalibrasi\s*:", "{{ job_info }} :", tr_clean)

        return f"{{% for {loop_var} %}}{tr_clean}{{% endfor %}}"

    xml = re.sub(r"<w:tr\b[^>]*>.*?</w:tr>", fix_tr, xml, flags=re.DOTALL)
    return xml


def prepare_docx_template_bytes(template_path: Path) -> io.BytesIO:
    """Preprocess a Word .docx template to clean split Jinja tags and return an

    in-memory stream ready for DocxTemplate.
    """
    template_path = Path(template_path)
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found at {template_path}")

    output_stream = io.BytesIO()
    with zipfile.ZipFile(template_path, "r") as zin:
        with zipfile.ZipFile(output_stream, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith(".xml"):
                    try:
                        text = data.decode("utf-8")
                        cleaned = clean_docx_xml_content(text)
                        data = cleaned.encode("utf-8")
                    except UnicodeDecodeError:
                        pass
                zout.writestr(item, data)
    output_stream.seek(0)
    return output_stream


# ==============================================================================
# 1. PENAWARAN (QUOTATION) GENERATOR (.docx)
# ==============================================================================


def generate_penawaran_docx(order: Order) -> bytes:
    """Generate official Penawaran Harga (.docx) using template_penawaran.docx."""
    tpl_path = TEMPLATES_DOCS_DIR / "template_penawaran.docx"
    stream = prepare_docx_template_bytes(tpl_path)
    doc = DocxTemplate(stream)

    items_data: list[dict[str, Any]] = []
    for item in order.job_items:
        desc = (
            item.specification_desc
            or f"{item.instrument_name} {item.brand or ''} {item.model or ''}".strip()
        )
        total_p = item.unit_price * item.qty
        items_data.append({
            "description": desc,
            "qty": item.qty,
            "unit": item.unit,
            "unit_price_fmt": format_currency_idr(item.unit_price),
            "total_price_fmt": format_currency_idr(total_p),
        })

    q_date = order.quotation_date or order.received_date or date.today()

    context = {
        "quotation_no": order.quotation_no or "-",
        "quotation_date": format_date_indo(q_date),
        "attachment_text": "1 Lampiran",
        "revision_no": order.revision_no or "00",
        "customer_name": order.customer.company_name,
        "customer_address": order.customer.address or "-",
        "pic_name": order.customer.pic_name or "-",
        "customer_email": order.customer.email or "-",
        "job_info": order.job_info or "Kalibrasi",
        "notes": order.notes or "",
        "items": items_data,
        "subtotal_fmt": format_currency_idr(order.subtotal),
        "ppn_fmt": format_currency_idr(order.ppn_amount),
        "grand_total_fmt": format_currency_idr(order.grand_total),
    }

    doc.render(context)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ==============================================================================
# 2. ORDERSHEET GENERATOR (.xlsx with drawings/logos preserved 100%)
# ==============================================================================


def generate_ordersheet_xlsx(order: Order) -> bytes:
    """Generate internal laboratory Ordersheet (.xlsx) preserving drawing1.xml

    (Text Box 2 company header and logos).
    """
    tpl_path = TEMPLATES_DOCS_DIR / "template_ordersheet.xlsx"
    if not tpl_path.exists():
        raise FileNotFoundError(f"Template not found at {tpl_path}")

    with zipfile.ZipFile(tpl_path, "r") as z_orig:
        orig_files = {name: z_orig.read(name) for name in z_orig.namelist()}

    wb = openpyxl.load_workbook(tpl_path)
    ws = wb["Order Sheet"]

    # Header section
    rec_date_str = (
        order.received_date.strftime("%d/%m/%Y") if order.received_date else "-"
    )
    ws["N9"] = rec_date_str
    ws["N10"] = order.po_no or "-"
    ws["N11"] = order.ordersheet_no or "-"

    # Items section (Rows 14..31)
    target_date_str = (
        order.target_sla_date.strftime("%d/%m/%Y") if order.target_sla_date else "-"
    )

    for idx, item in enumerate(order.job_items[:18]):
        row_num = 14 + idx
        ws[f"A{row_num}"] = idx + 1
        ws[f"B{row_num}"] = item.instrument_name
        ws[f"E{row_num}"] = item.qty
        ws[f"F{row_num}"] = item.certificate_no or "-"
        ws[f"G{row_num}"] = item.serial_number
        ws[f"I{row_num}"] = rec_date_str
        ws[f"K{row_num}"] = target_date_str
        ws[f"M{row_num}"] = item.accessories or "Unit Only"

    # Customer & address lines (A33..A36)
    ws["A33"] = order.customer.company_name
    lines = [
        line.strip()
        for line in (order.customer.address or "").split("\n")
        if line.strip()
    ]
    ws["A34"] = lines[0] if len(lines) > 0 else "-"
    ws["A35"] = lines[1] if len(lines) > 1 else ""
    ws["A36"] = (
        f"Attn: {order.customer.pic_name} ({order.customer.phone})"
        if order.customer.pic_name
        else ""
    )

    raw_stream = io.BytesIO()
    wb.save(raw_stream)
    raw_stream.seek(0)

    # Merge saved openpyxl workbook with original drawing1.xml & media images
    final_stream = io.BytesIO()
    with zipfile.ZipFile(raw_stream, "r") as zin:
        saved_files = {item.filename: zin.read(item.filename) for item in zin.infolist()}

    # 1. Update sheet1.xml to link drawing
    sheet1_xml = saved_files["xl/worksheets/sheet1.xml"].decode("utf-8")
    if "xmlns:r=" not in sheet1_xml:
        sheet1_xml = sheet1_xml.replace(
            "<worksheet ",
            '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ',
        )
    if "<drawing" not in sheet1_xml:
        sheet1_xml = sheet1_xml.replace(
            "</worksheet>", '<drawing r:id="rId2"/></worksheet>'
        )
    saved_files["xl/worksheets/sheet1.xml"] = sheet1_xml.encode("utf-8")

    # 2. Update [Content_Types].xml
    ct_xml = saved_files["[Content_Types].xml"].decode("utf-8")
    if "/xl/drawings/drawing1.xml" not in ct_xml:
        ct_xml = ct_xml.replace(
            "</Types>",
            '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
            '<Default Extension="png" ContentType="image/png"/>'
            '<Default Extension="jpeg" ContentType="image/jpeg"/>'
            "</Types>",
        )
    saved_files["[Content_Types].xml"] = ct_xml.encode("utf-8")

    # 3. Restore drawings and media
    for name in [
        "xl/drawings/drawing1.xml",
        "xl/drawings/_rels/drawing1.xml.rels",
        "xl/worksheets/_rels/sheet1.xml.rels",
        "xl/media/image1.png",
        "xl/media/image2.jpeg",
        "xl/media/image3.jpeg",
    ]:
        if name in orig_files:
            saved_files[name] = orig_files[name]

    with zipfile.ZipFile(final_stream, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, content in saved_files.items():
            zout.writestr(name, content)

    final_stream.seek(0)
    return final_stream.getvalue()


# ==============================================================================
# 3. TANDA TERIMA (TT) PENGEMBALIAN ALAT (.docx)
# ==============================================================================


def generate_tt_docx(
    order: Order, recipient_name: str | None = None, keterangan: str = "Barang"
) -> bytes:
    """Generate Tanda Terima Pengembalian Barang (.docx) using template_tt.docx."""
    tpl_path = TEMPLATES_DOCS_DIR / "template_tt.docx"
    stream = prepare_docx_template_bytes(tpl_path)
    doc = DocxTemplate(stream)

    items_data: list[dict[str, Any]] = []
    for item in order.job_items:
        items_data.append({
            "instrument_name": item.instrument_name,
            "serial_or_owner": item.serial_number,
            "qty": item.qty,
            "keterangan": keterangan,
        })

    ret_date = order.returned_date or date.today()
    recip = (
        recipient_name
        or (order.customer.company_name if order.customer else None)
        or "Customer"
    )

    context = {
        "tt_no": order.tt_no or "-",
        "hari_tanggal": format_day_date_indo(ret_date),
        "recipient_name": recip,
        "po_no": order.po_no or "-",
        "items": items_data,
    }

    doc.render(context)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ==============================================================================
# 4. BERITA ACARA SERAH TERIMA PEKERJAAN (BA) (.docx)
# ==============================================================================


def generate_ba_docx(order: Order) -> bytes:
    """Generate Berita Acara Penyelesaian Kalibrasi (.docx) using template_ba.docx.

    Guardrail:
    Raises HTTP 403 Forbidden unless order.po_status == POStatus.RECEIVED
    and all items are CERT_READY or DELIVERED.
    """
    if order.po_status != POStatus.RECEIVED:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Berita Acara tidak dapat dicetak: Status PO customer masih PENDING. "
                "Harus berstatus RECEIVED sebelum Berita Acara diterbitkan."
            ),
        )

    for item in order.job_items:
        if item.status not in (JobStatus.CERT_READY, JobStatus.DELIVERED):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=(
                    f"Berita Acara tidak dapat dicetak: Instrumen '{item.instrument_name}' "
                    f"masih dalam status '{item.status.value}'. Semua alat harus telah berstatus "
                    "CERT_READY atau DELIVERED."
                ),
            )

    tpl_path = TEMPLATES_DOCS_DIR / "template_ba.docx"
    stream = prepare_docx_template_bytes(tpl_path)
    doc = DocxTemplate(stream)

    items_data: list[dict[str, Any]] = []
    for item in order.job_items:
        desc = (
            item.specification_desc
            or f"{item.instrument_name} {item.brand or ''} {item.model or ''}".strip()
        )
        items_data.append({
            "description": desc,
            "qty": item.qty,
            "unit": item.unit,
        })

    today = order.returned_date or date.today()

    context = {
        "bast_no": order.bast_no or "-",
        "hari": format_day_indo(today),
        "tanggal_indo": format_date_indo(today),
        "job_info": order.job_info or "Kalibrasi",
        "customer_name": order.customer.company_name,
        "customer_address": order.customer.address or "-",
        "po_no": order.po_no or "-",
        "items": items_data,
    }

    doc.render(context)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ==============================================================================
# 5. PO KELUAR SUBCONTRACTOR (.docx)
# ==============================================================================


def generate_po_keluar_docx(po_out: PurchaseOrderOut) -> bytes:
    """Generate PO Keluar Subcontractor (.docx) using template_po_keluar.docx."""
    tpl_path = TEMPLATES_DOCS_DIR / "template_po_keluar.docx"
    stream = prepare_docx_template_bytes(tpl_path)
    doc = DocxTemplate(stream)

    items_data: list[dict[str, Any]] = []
    # If po_out has job_items linked, use them; otherwise, pull all items of that order
    items_to_render = po_out.job_items
    if not items_to_render:
        items_to_render = [
            it
            for it in po_out.order.job_items
            if it.subcontractor_id == po_out.subcontractor_id
        ]
    if not items_to_render:
        items_to_render = po_out.order.job_items

    for item in items_to_render:
        name_str = f"{item.instrument_name} ({item.brand or ''} {item.model or ''}, SN: {item.serial_number})".strip()
        unit_p = item.subcon_price if item.subcon_price > 0 else item.unit_price
        total_p = unit_p * item.qty
        items_data.append({
            "item_name": name_str,
            "qty": item.qty,
            "unit": item.unit,
            "unit_price_fmt": format_currency_idr(unit_p),
            "total_price_fmt": format_currency_idr(total_p),
        })

    # Onsite fee if any
    if po_out.onsite_fee > 0:
        items_data.append({
            "item_name": "Biaya Onsite / Akomodasi",
            "qty": 1,
            "unit": "Lot",
            "unit_price_fmt": format_currency_idr(po_out.onsite_fee),
            "total_price_fmt": format_currency_idr(po_out.onsite_fee),
        })

    context = {
        "vendor_name": po_out.subcontractor.vendor_name,
        "vendor_address": po_out.subcontractor.address or "-",
        "vendor_phone": po_out.subcontractor.phone or "-",
        "vendor_pic": po_out.subcontractor.pic_name or "-",
        "po_out_no": po_out.po_out_no,
        "po_date": format_date_indo(po_out.po_date),
        "delivery_date": format_date_indo(po_out.delivery_date) if po_out.delivery_date else "-",
        "vendor_quotation_no": po_out.vendor_quotation_no or "-",
        "job_info": po_out.job_info or "Kalibrasi",
        "customer_name": po_out.order.customer.company_name,
        "customer_address": po_out.order.customer.address or "-",
        "notes": po_out.notes or "",
        "items": items_data,
        "subtotal_fmt": format_currency_idr(po_out.subtotal),
        "ppn_fmt": format_currency_idr(po_out.ppn_amount),
        "grand_total_fmt": format_currency_idr(po_out.grand_total),
    }

    doc.render(context)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ==============================================================================
# 6. INVOICE GENERATOR (.docx)
# ==============================================================================


def generate_invoice_docx(order: Order) -> bytes:
    """Generate Invoice (.docx) using template_invoice.docx.

    Guardrail:
    Raises HTTP 403 Forbidden unless order.bast_no is issued.
    """
    if not order.bast_no or not order.bast_no.strip():
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Invoice tidak dapat diterbitkan: Berita Acara Serah Terima (BAST) belum dibuat. "
                "Selesaikan penyerahan alat dan terbitkan BAST terlebih dahulu."
            ),
        )

    tpl_path = TEMPLATES_DOCS_DIR / "template_invoice.docx"
    stream = prepare_docx_template_bytes(tpl_path)
    doc = DocxTemplate(stream)

    items_data: list[dict[str, Any]] = []
    for item in order.job_items:
        desc = (
            item.specification_desc
            or f"{item.instrument_name} {item.brand or ''} {item.model or ''}".strip()
        )
        total_p = item.unit_price * item.qty
        items_data.append({
            "description": desc,
            "qty": item.qty,
            "unit": item.unit,
            "unit_price_fmt": format_currency_idr(item.unit_price),
            "total_price_fmt": format_currency_idr(total_p),
        })

    inv_date = order.invoice_date or date.today()

    context = {
        "customer_name": order.customer.company_name,
        "customer_address": order.customer.address or "-",
        "invoice_date": format_date_indo(inv_date),
        "po_no": order.po_no or "-",
        "invoice_no": order.invoice_no or "-",
        "terms_pembayaran": order.terms_pembayaran or "30 Hari",
        "job_info": order.job_info or "Kalibrasi",
        "items": items_data,
        "subtotal_fmt": format_currency_idr(order.subtotal),
        "ppn_fmt": format_currency_idr(order.ppn_amount),
        "grand_total_fmt": format_currency_idr(order.grand_total),
    }

    doc.render(context)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
