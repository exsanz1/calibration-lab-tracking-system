"""Tests for 7-Stage Operational Workflow, Subcontractors, and Template Document Generation."""

from __future__ import annotations

import io
import zipfile
from datetime import date, timedelta

import openpyxl  # type: ignore[import-untyped]
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.customer import Customer
from app.models.job_item import JobItem, JobStatus, Scope
from app.models.order import InvoiceStatus, Order, POStatus, ServiceType
from app.models.po_out import POOutStatus, PurchaseOrderOut
from app.models.subcontractor import Subcontractor
from app.services.doc_generator import (
    generate_ba_docx,
    generate_invoice_docx,
    generate_ordersheet_xlsx,
    generate_penawaran_docx,
    generate_po_keluar_docx,
    generate_tt_docx,
)
from app.services.number_generator import (
    generate_bast_no,
    generate_certificate_no,
    generate_invoice_no,
    generate_ordersheet_no,
    generate_po_out_no,
    generate_quotation_no,
    generate_tt_no,
)


@pytest.fixture
def sample_customer(db_session: Session) -> Customer:
    c = Customer(
        company_name="PT Multi Integra Test",
        pic_name="Ahmad Fauzi",
        phone="021-1234567",
        email="fauzi@multiintegra.com",
        address="Kawasan Industri Pulogadung Blok B No. 8, Jakarta Timur",
    )
    db_session.add(c)
    db_session.commit()
    db_session.refresh(c)
    return c


@pytest.fixture
def sample_subcontractor(db_session: Session) -> Subcontractor:
    sub = Subcontractor(
        vendor_name="PT Kalibrasi Mandiri Utama",
        pic_name="Hendra Saputra",
        phone="021-888999",
        email="hendra@kmu-lab.co.id",
        address="Ruko Grand Wisata Blok AA No. 12, Bekasi",
    )
    db_session.add(sub)
    db_session.commit()
    db_session.refresh(sub)
    return sub


@pytest.fixture
def sample_order(db_session: Session, sample_customer: Customer) -> Order:
    today = date.today()
    order = Order(
        customer_id=sample_customer.id,
        job_info="Kalibrasi",
        quotation_no="0211/PNW-NSI/X/2026",
        quotation_date=today - timedelta(days=10),
        revision_no="00",
        notes="Sertifikat rangkap 1",
        po_no="PO-TEST-001",
        po_status=POStatus.RECEIVED,
        service_type=ServiceType.IN_HOUSE,
        ordersheet_no="K.0156.1026",
        received_date=today - timedelta(days=5),
        target_sla_date=today + timedelta(days=5),
        tt_no="0285/TT-NSI/IX/26",
        returned_date=today,
        bast_no="0189/BAP-NSI/IX/2026",
        invoice_no="0182/INV-NSI/IX/2026",
        invoice_date=today,
        terms_pembayaran="30 Hari",
        invoice_status=InvoiceStatus.INVOICED,
        subtotal=0,
        ppn_percent=11,
        ppn_amount=0,
        grand_total=0,
    )
    db_session.add(order)
    db_session.commit()
    db_session.refresh(order)

    # Add 2 items
    item1 = JobItem(
        order_id=order.id,
        qr_tag="QR-T-001",
        instrument_name="Digital Pressure Gauge",
        specification_desc="Digital Pressure Gauge 0-100 bar, Conn 1/2 NPT",
        brand="WIKA",
        model="CPG1500",
        serial_number="SN-DPG-001",
        scope=Scope.PRESSURE,
        status=JobStatus.CERT_READY,
        certificate_no="NSI.0478.1026",
        qty=2,
        unit="Unit",
        unit_price=500000,
        release_allowed=True,
    )
    item2 = JobItem(
        order_id=order.id,
        qr_tag="QR-T-002",
        instrument_name="Temperature Calibrator",
        specification_desc="Portable Dry Block Calibrator Range 33 to 300 C",
        brand="Fluke",
        model="9100S",
        serial_number="SN-TC-002",
        scope=Scope.TEMP,
        status=JobStatus.CERT_READY,
        certificate_no="NSI.0479.1026",
        qty=1,
        unit="Unit",
        unit_price=1200000,
        release_allowed=True,
    )
    db_session.add_all([item1, item2])
    db_session.commit()

    # Recalculate financials
    subtot = (500000 * 2) + (1200000 * 1)  # 2,200,000
    ppn = int(subtot * 0.11)  # 242,000
    order.subtotal = subtot
    order.ppn_amount = ppn
    order.grand_total = subtot + ppn
    db_session.commit()
    db_session.refresh(order)

    return order


# ==========================================
# 1. Number Generator Unit Tests
# ==========================================


def test_number_generation_formats() -> None:
    d = date(2026, 10, 15)
    assert generate_quotation_no(42, d) == "0042/PNW-NSI/X/2026"
    assert generate_ordersheet_no(156, d) == "K.0156.1026"
    assert generate_bast_no(189, d) == "0189/BAP-NSI/X/2026"
    assert generate_invoice_no(182, d) == "0182/INV-NSI/X/2026"
    assert generate_tt_no(285, d) == "0285/TT-NSI/X/26"
    assert generate_po_out_no(70, d) == "070/PO-NSI/X/2026"
    assert generate_certificate_no(478, d) == "NSI.0478.1026"


# ==========================================
# 2. Document Generator Functions (Docx & Excel)
# ==========================================


def test_generate_penawaran_docx(sample_order: Order) -> None:
    doc_bytes = generate_penawaran_docx(sample_order)
    assert len(doc_bytes) > 1000
    # Must be valid zip archive
    zf = zipfile.ZipFile(io.BytesIO(doc_bytes))
    assert "word/document.xml" in zf.namelist()


def test_generate_ordersheet_xlsx(sample_order: Order) -> None:
    xlsx_bytes = generate_ordersheet_xlsx(sample_order)
    assert len(xlsx_bytes) > 1000

    # Verify openpyxl can load it and data is present
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Order Sheet"]
    assert ws["N11"].value == sample_order.ordersheet_no
    assert ws["B14"].value == sample_order.job_items[0].instrument_name

    # Verify drawing1.xml was preserved
    zf = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    assert "xl/drawings/drawing1.xml" in zf.namelist()
    assert "xl/drawings/_rels/drawing1.xml.rels" in zf.namelist()


def test_generate_tt_docx(sample_order: Order) -> None:
    doc_bytes = generate_tt_docx(sample_order)
    assert len(doc_bytes) > 1000
    zf = zipfile.ZipFile(io.BytesIO(doc_bytes))
    assert "word/document.xml" in zf.namelist()


def test_generate_ba_docx(sample_order: Order) -> None:
    doc_bytes = generate_ba_docx(sample_order)
    assert len(doc_bytes) > 1000
    zf = zipfile.ZipFile(io.BytesIO(doc_bytes))
    assert "word/document.xml" in zf.namelist()


def test_generate_invoice_docx(sample_order: Order) -> None:
    doc_bytes = generate_invoice_docx(sample_order)
    assert len(doc_bytes) > 1000
    zf = zipfile.ZipFile(io.BytesIO(doc_bytes))
    assert "word/document.xml" in zf.namelist()


def test_generate_po_keluar_docx(
    db_session: Session, sample_order: Order, sample_subcontractor: Subcontractor
) -> None:
    po_out = PurchaseOrderOut(
        po_out_no="0070/PO-NSI/X/2026",
        subcontractor_id=sample_subcontractor.id,
        order_id=sample_order.id,
        job_info="Kalibrasi",
        po_date=date.today(),
        delivery_date=date.today() + timedelta(days=7),
        vendor_quotation_no="QUO-TEST-123",
        notes="Sertifikat harap dikirim softcopy",
        onsite_fee=0,
        subtotal=1000000,
        ppn_amount=110000,
        grand_total=1110000,
        status=POOutStatus.ISSUED,
    )
    db_session.add(po_out)
    db_session.commit()
    db_session.refresh(po_out)

    doc_bytes = generate_po_keluar_docx(po_out)
    assert len(doc_bytes) > 1000
    zf = zipfile.ZipFile(io.BytesIO(doc_bytes))
    assert "word/document.xml" in zf.namelist()


# ==========================================
# 3. Guardrail Unit Tests (403 Forbidden)
# ==========================================


def test_ba_guardrail_forbidden_when_po_pending(
    db_session: Session, sample_order: Order
) -> None:
    sample_order.po_status = POStatus.PENDING
    db_session.commit()

    with pytest.raises(Exception) as excinfo:
        generate_ba_docx(sample_order)
    assert "403" in str(excinfo.value) or "PO belum diterima" in str(excinfo.value)


def test_ba_guardrail_forbidden_when_items_not_ready(
    db_session: Session, sample_order: Order
) -> None:
    sample_order.po_status = POStatus.RECEIVED
    sample_order.job_items[0].status = JobStatus.CALIBRATING
    db_session.commit()

    with pytest.raises(Exception) as excinfo:
        generate_ba_docx(sample_order)
    assert "403" in str(excinfo.value) or "belum selesai" in str(excinfo.value)


def test_invoice_guardrail_forbidden_without_bast(
    db_session: Session, sample_order: Order
) -> None:
    sample_order.bast_no = None
    db_session.commit()

    with pytest.raises(Exception) as excinfo:
        generate_invoice_docx(sample_order)
    assert "403" in str(excinfo.value) or "BAST belum dibuat" in str(excinfo.value)


# ==========================================
# 4. HTTP Router & Action Endpoints
# ==========================================


def test_http_download_docs(client: TestClient, sample_order: Order) -> None:
    res = client.get(f"/orders/{sample_order.id}/docs/penawaran")
    assert res.status_code == 200
    assert "application/vnd.openxmlformats" in res.headers["content-type"]

    res = client.get(f"/orders/{sample_order.id}/docs/ordersheet")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]

    res = client.get(f"/orders/{sample_order.id}/docs/tt")
    assert res.status_code == 200

    res = client.get(f"/orders/{sample_order.id}/docs/ba")
    assert res.status_code == 200

    res = client.get(f"/orders/{sample_order.id}/docs/invoice")
    assert res.status_code == 200


def test_http_order_detail_view(client: TestClient, sample_order: Order) -> None:
    res = client.get(f"/orders/{sample_order.id}")
    assert res.status_code == 200
    assert "Alur Kerja Operasional 7 Tahap" in res.text
    assert sample_order.quotation_no in res.text
    assert "Unduh Penawaran" in res.text


def test_http_quick_action_automations(
    client: TestClient, db_session: Session, sample_customer: Customer
) -> None:
    # Create blank order
    order = Order(
        customer_id=sample_customer.id,
        service_type=ServiceType.IN_HOUSE,
        po_status=POStatus.RECEIVED,
    )
    db_session.add(order)
    db_session.commit()
    db_session.refresh(order)

    # 1. Action gen quotation
    res = client.post(
        f"/orders/{order.id}/actions/gen-quotation", follow_redirects=True
    )
    assert res.status_code == 200
    db_session.refresh(order)
    assert order.quotation_no is not None
    assert "/PNW-NSI/" in order.quotation_no

    # 2. Action gen ordersheet
    res = client.post(
        f"/orders/{order.id}/actions/gen-ordersheet", follow_redirects=True
    )
    assert res.status_code == 200
    db_session.refresh(order)
    assert order.ordersheet_no is not None
    assert order.ordersheet_no.startswith("K.")

    # 3. Action gen BAST without ready items should redirect with error
    res = client.post(f"/orders/{order.id}/actions/gen-bast-tt", follow_redirects=True)
    assert res.status_code == 200
    # Must fail because no items are CERT_READY
    assert "error=" in str(res.url) or "tidak dapat diterbitkan" in res.text

    # Now add a ready item
    it = JobItem(
        order_id=order.id,
        qr_tag="QR-ACTION-01",
        instrument_name="Test Tool",
        serial_number="SN-01",
        scope=Scope.PRESSURE,
        status=JobStatus.CERT_READY,
        release_allowed=True,
    )
    db_session.add(it)
    db_session.commit()

    # Now BAST & TT should succeed
    res = client.post(f"/orders/{order.id}/actions/gen-bast-tt", follow_redirects=True)
    assert res.status_code == 200
    db_session.refresh(order)
    assert order.bast_no is not None
    assert "/BAP-NSI/" in order.bast_no
    assert order.tt_no is not None

    # 4. Action gen invoice
    res = client.post(f"/orders/{order.id}/actions/gen-invoice", follow_redirects=True)
    assert res.status_code == 200
    db_session.refresh(order)
    assert order.invoice_no is not None
    assert "/INV-NSI/" in order.invoice_no
    assert order.invoice_status == InvoiceStatus.INVOICED


# ==========================================
# 5. Subcontractor & PO Keluar CRUD
# ==========================================


def test_subcontractor_and_po_out_crud(
    client: TestClient, sample_order: Order
) -> None:
    # 1. Create subcontractor via API
    sub_payload = {
        "vendor_name": "PT Akurasi Instrumen Nusantara",
        "pic_name": "Bambang Sudiro",
        "phone": "021-998877",
        "email": "bambang@akurasi.com",
        "address": "Kawasan Jababeka, Cikarang",
    }
    res = client.post("/api/subcontractors", json=sub_payload)
    assert res.status_code == 201
    sub_id = res.json()["id"]

    # 2. View subcontractor HTML list and edit
    res = client.get("/subcontractors")
    assert res.status_code == 200
    assert "PT Akurasi Instrumen Nusantara" in res.text

    res = client.get(f"/subcontractors/{sub_id}/edit")
    assert res.status_code == 200

    # 3. Create PO Keluar
    po_payload = {
        "po_out_no": "0099/PO-NSI/X/2026",
        "subcontractor_id": sub_id,
        "order_id": sample_order.id,
        "job_info": "Kalibrasi Khusus",
        "po_date": str(date.today()),
        "vendor_quotation_no": "Q-VEND-99",
        "notes": "Kalibrasi lab 17025",
        "onsite_fee": 150000,
        "subtotal": 1500000,
        "ppn_amount": 165000,
        "grand_total": 1815000,
        "status": "ISSUED",
    }
    res = client.post("/api/po-out", json=po_payload)
    assert res.status_code == 201
    po_id = res.json()["id"]

    # 4. View PO Keluar detail HTML
    res = client.get(f"/po-out/{po_id}")
    assert res.status_code == 200
    assert "0099/PO-NSI/X/2026" in res.text

    # 5. Download PO Keluar Word document
    res = client.get(f"/po-out/{po_id}/doc")
    assert res.status_code == 200
    assert len(res.content) > 1000
